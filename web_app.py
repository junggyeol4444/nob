"""NovelWriter Local — Flask 웹 애플리케이션"""
import os
import re
import io
import sys
import json
import time
import threading
from datetime import datetime

sys.path.insert(0, os.path.dirname(__file__))

from flask import Flask, request, jsonify, render_template, send_file, abort
from flask_socketio import SocketIO, emit

from src.database.db_manager import DatabaseManager
from src.database.models import Project, Chapter, Character, WorldBuilding, Outline, GenerationJob
from src.ai.llm_client import LLMClient
from src.ai.context_manager import ContextManager
from src.ai.auto_generator import AutoGenerator, cancel_job, pause_job, resume_job
from src.utils import config

# ── 앱 초기화 ──────────────────────────────────────────────────────────────

app = Flask(__name__, template_folder="templates")
app.config["SECRET_KEY"] = os.urandom(24)
socketio = SocketIO(app, cors_allowed_origins="*", async_mode="threading")

cfg = config.load_config()
db_path = cfg.get("db_path") or config.get_default_db_path()
db = DatabaseManager(db_path)


def _make_llm_client():
    c = config.load_config()
    return LLMClient(
        endpoint=c.get("llm_endpoint", "http://localhost:11434"),
        model=c.get("llm_model", "mistral"),
        temperature=c.get("temperature", 0.7),
        max_tokens=c.get("max_tokens", 4096),
    )


def _make_context_manager():
    return ContextManager(db, _make_llm_client())


def _on_job_progress(job_id, job_dict):
    socketio.emit("generation_progress", {"job_id": job_id, "job": job_dict})


# ── 메인 페이지 ────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


# ── 설정 API ───────────────────────────────────────────────────────────────

@app.route("/api/settings", methods=["GET"])
def get_settings():
    c = config.load_config()
    c.pop("db_path", None)
    return jsonify(c)


@app.route("/api/settings", methods=["PUT"])
def update_settings():
    data = request.json or {}
    c = config.load_config()
    for key in ("llm_endpoint", "llm_model", "temperature", "max_tokens",
                "auto_save_interval", "default_word_count"):
        if key in data:
            c[key] = data[key]
    config.save_config(c)
    return jsonify({"ok": True})


@app.route("/api/llm/check", methods=["GET"])
def check_llm():
    client = _make_llm_client()
    reachable = client.is_reachable()
    models = client.list_models() if reachable else []
    return jsonify({"reachable": reachable, "models": models})


# ── 프로젝트 API ───────────────────────────────────────────────────────────

@app.route("/api/projects", methods=["GET"])
def list_projects():
    projects = db.get_all_projects()
    result = []
    for p in projects:
        d = p.to_dict()
        chapters = db.get_chapters_by_project(p.id)
        d["chapter_count"] = len(chapters)
        d["generated_count"] = sum(1 for c in chapters if c.status == "generated")
        result.append(d)
    return jsonify(result)


@app.route("/api/projects", methods=["POST"])
def create_project():
    data = request.json or {}
    project = Project(
        title=data.get("title", "새 프로젝트"),
        genre=data.get("genre", "판타지"),
        description=data.get("description", ""),
        tone=data.get("tone", ""),
        viewpoint=data.get("viewpoint", "3인칭"),
        plot=data.get("plot", ""),
        plot_source=data.get("plot_source", "user"),
        total_chapters=int(data.get("total_chapters", 100)),
        chapter_length=int(data.get("chapter_length", 2000)),
        status=data.get("status", "active"),
    )
    project.id = db.create_project(project)
    return jsonify(project.to_dict()), 201


@app.route("/api/projects/<int:project_id>", methods=["GET"])
def get_project(project_id):
    p = db.get_project(project_id)
    if not p:
        abort(404)
    return jsonify(p.to_dict())


@app.route("/api/projects/<int:project_id>", methods=["PUT"])
def update_project(project_id):
    p = db.get_project(project_id)
    if not p:
        abort(404)
    data = request.json or {}
    p.title = data.get("title", p.title)
    p.genre = data.get("genre", p.genre)
    p.description = data.get("description", p.description)
    p.tone = data.get("tone", p.tone)
    p.viewpoint = data.get("viewpoint", p.viewpoint)
    p.plot = data.get("plot", p.plot)
    p.plot_source = data.get("plot_source", p.plot_source)
    p.total_chapters = int(data.get("total_chapters", p.total_chapters))
    p.chapter_length = int(data.get("chapter_length", p.chapter_length))
    p.status = data.get("status", p.status)
    db.update_project(p)
    return jsonify(p.to_dict())


@app.route("/api/projects/<int:project_id>", methods=["DELETE"])
def delete_project(project_id):
    db.delete_project(project_id)
    return jsonify({"ok": True})


# ── 목차 API ───────────────────────────────────────────────────────────────

@app.route("/api/projects/<int:project_id>/outlines", methods=["GET"])
def list_outlines(project_id):
    outlines = db.get_outlines_by_project(project_id)
    return jsonify([o.to_dict() for o in outlines])


@app.route("/api/projects/<int:project_id>/outlines", methods=["POST"])
def create_outline(project_id):
    data = request.json or {}
    outline = Outline(
        project_id=project_id,
        chapter_num=int(data.get("chapter_num", 1)),
        title=data.get("title", ""),
        description=data.get("description", ""),
        source=data.get("source", "user"),
        order_index=data.get("order_index"),
    )
    outline.id = db.create_outline(outline)
    return jsonify(outline.to_dict()), 201


@app.route("/api/projects/<int:project_id>/outlines/<int:outline_id>", methods=["PUT"])
def update_outline(project_id, outline_id):
    outline = db.get_outline(outline_id)
    if not outline:
        abort(404)
    data = request.json or {}
    outline.chapter_num = int(data.get("chapter_num", outline.chapter_num))
    outline.title = data.get("title", outline.title)
    outline.description = data.get("description", outline.description)
    outline.source = data.get("source", outline.source)
    outline.order_index = data.get("order_index", outline.order_index)
    db.update_outline(outline)
    return jsonify(outline.to_dict())


@app.route("/api/projects/<int:project_id>/outlines/<int:outline_id>", methods=["DELETE"])
def delete_outline(project_id, outline_id):
    db.delete_outline(outline_id)
    return jsonify({"ok": True})


@app.route("/api/projects/<int:project_id>/outlines/generate", methods=["POST"])
def generate_outlines(project_id):
    """AI 목차 자동 생성"""
    p = db.get_project(project_id)
    if not p:
        abort(404)
    try:
        ctx = _make_context_manager()
        raw = ctx.generate_outline(project_id)
        # 파싱: "숫자. 제목" 또는 "숫자편: 제목" 형식
        outlines = _parse_outline_text(raw, project_id, p.total_chapters)
        # 기존 목차 삭제 후 새로 저장
        db.delete_all_outlines(project_id)
        for o in outlines:
            db.create_outline(o)
        saved = db.get_outlines_by_project(project_id)
        return jsonify([o.to_dict() for o in saved])
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/projects/<int:project_id>/outlines/bulk", methods=["POST"])
def bulk_save_outlines(project_id):
    """목차 일괄 저장"""
    data = request.json or {}
    items = data.get("outlines", [])
    db.delete_all_outlines(project_id)
    for item in items:
        outline = Outline(
            project_id=project_id,
            chapter_num=int(item.get("chapter_num", 1)),
            title=item.get("title", ""),
            description=item.get("description", ""),
            source=item.get("source", "user"),
            order_index=item.get("order_index"),
        )
        db.create_outline(outline)
    saved = db.get_outlines_by_project(project_id)
    return jsonify([o.to_dict() for o in saved])


def _parse_outline_text(text: str, project_id: int, total: int) -> list:
    """AI 생성 목차 텍스트를 Outline 객체 목록으로 파싱"""
    outlines = []
    lines = text.strip().split("\n")
    seen = set()
    for line in lines:
        line = line.strip()
        if not line:
            continue
        # 패턴: "1. 제목", "1편: 제목", "1화: 제목", "제1편 제목"
        m = re.match(r"^(?:제\s*)?(\d+)[편화\.][\s:：.]*(.+)$", line)
        if not m:
            m = re.match(r"^(\d+)\s*[.)]\s*(.+)$", line)
        if m:
            num = int(m.group(1))
            title = m.group(2).strip()
            if 1 <= num <= total and num not in seen:
                seen.add(num)
                outlines.append(Outline(
                    project_id=project_id,
                    chapter_num=num,
                    title=title,
                    source="ai",
                    order_index=num,
                ))
    return sorted(outlines, key=lambda o: o.chapter_num)


# ── 챕터 API ───────────────────────────────────────────────────────────────

@app.route("/api/projects/<int:project_id>/chapters", methods=["GET"])
def list_chapters(project_id):
    chapters = db.get_chapters_by_project(project_id)
    result = []
    for c in chapters:
        d = c.to_dict()
        d.pop("content", None)  # 목록에서는 본문 제외 (용량 절약)
        result.append(d)
    return jsonify(result)


@app.route("/api/projects/<int:project_id>/chapters", methods=["POST"])
def create_chapter(project_id):
    data = request.json or {}
    chapter = Chapter(
        project_id=project_id,
        outline_id=data.get("outline_id"),
        chapter_num=int(data.get("chapter_num", db.get_next_chapter_num(project_id))),
        title=data.get("title", ""),
        content=data.get("content", ""),
        summary=data.get("summary", ""),
        word_count=len(data.get("content", "")),
        status=data.get("status", "draft"),
    )
    chapter.id = db.create_chapter(chapter)
    return jsonify(chapter.to_dict()), 201


@app.route("/api/projects/<int:project_id>/chapters/<int:chapter_id>", methods=["GET"])
def get_chapter(project_id, chapter_id):
    ch = db.get_chapter(chapter_id)
    if not ch or ch.project_id != project_id:
        abort(404)
    return jsonify(ch.to_dict())


@app.route("/api/projects/<int:project_id>/chapters/<int:chapter_id>", methods=["PUT"])
def update_chapter(project_id, chapter_id):
    ch = db.get_chapter(chapter_id)
    if not ch or ch.project_id != project_id:
        abort(404)
    data = request.json or {}
    if "content" in data:
        ch.content = data["content"]
        ch.word_count = len(data["content"])
    ch.title = data.get("title", ch.title)
    ch.summary = data.get("summary", ch.summary)
    ch.status = data.get("status", ch.status)
    db.update_chapter(ch)
    return jsonify(ch.to_dict())


@app.route("/api/projects/<int:project_id>/chapters/<int:chapter_id>", methods=["DELETE"])
def delete_chapter(project_id, chapter_id):
    db.delete_chapter(chapter_id)
    return jsonify({"ok": True})


@app.route("/api/projects/<int:project_id>/chapters/<int:chapter_id>/history", methods=["GET"])
def get_chapter_history(project_id, chapter_id):
    history = db.get_generation_history(chapter_id)
    return jsonify([h.to_dict() for h in history])


# ── 캐릭터 API ─────────────────────────────────────────────────────────────

@app.route("/api/projects/<int:project_id>/characters", methods=["GET"])
def list_characters(project_id):
    chars = db.get_characters_by_project(project_id)
    return jsonify([c.to_dict() for c in chars])


@app.route("/api/projects/<int:project_id>/characters", methods=["POST"])
def create_character(project_id):
    data = request.json or {}
    char = Character(
        project_id=project_id,
        name=data.get("name", ""),
        role=data.get("role", ""),
        age=data.get("age"),
        gender=data.get("gender", ""),
        personality=data.get("personality", ""),
        appearance=data.get("appearance", ""),
        background=data.get("background", ""),
        abilities=data.get("abilities", ""),
        relationships=data.get("relationships", ""),
        character_arc=data.get("character_arc", ""),
        notes=data.get("notes", ""),
        source=data.get("source", "user"),
    )
    char.id = db.create_character(char)
    return jsonify(char.to_dict()), 201


@app.route("/api/projects/<int:project_id>/characters/<int:char_id>", methods=["PUT"])
def update_character(project_id, char_id):
    char = db.get_character(char_id)
    if not char or char.project_id != project_id:
        abort(404)
    data = request.json or {}
    for field in ("name", "role", "gender", "personality", "appearance",
                  "background", "abilities", "relationships", "character_arc", "notes", "source"):
        if field in data:
            setattr(char, field, data[field])
    if "age" in data:
        char.age = data["age"]
    db.update_character(char)
    return jsonify(char.to_dict())


@app.route("/api/projects/<int:project_id>/characters/<int:char_id>", methods=["DELETE"])
def delete_character(project_id, char_id):
    db.delete_character(char_id)
    return jsonify({"ok": True})


@app.route("/api/projects/<int:project_id>/characters/generate", methods=["POST"])
def generate_character(project_id):
    p = db.get_project(project_id)
    if not p:
        abort(404)
    data = request.json or {}
    role = data.get("role", "주인공")
    try:
        ctx = _make_context_manager()
        raw = ctx.generate_character(project_id, role)
        return jsonify({"role": role, "raw": raw, "parsed": _parse_character_text(raw, project_id, role)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def _parse_character_text(text: str, project_id: int, role: str) -> dict:
    """AI 생성 캐릭터 텍스트 파싱"""
    fields = {"name": "", "age": None, "gender": "", "personality": "",
              "appearance": "", "background": "", "abilities": "", "character_arc": ""}
    mapping = {
        "이름": "name", "나이": "age", "성별": "gender", "성격": "personality",
        "외모": "appearance", "배경": "background", "능력": "abilities",
        "캐릭터 아크": "character_arc",
    }
    for line in text.split("\n"):
        for label, key in mapping.items():
            if line.startswith(label + ":") or line.startswith(label + "："):
                val = line.split(":", 1)[-1].strip().lstrip("：").strip()
                if key == "age":
                    m = re.search(r"\d+", val)
                    fields[key] = int(m.group()) if m else None
                else:
                    fields[key] = val
    fields["project_id"] = project_id
    fields["role"] = role
    fields["source"] = "ai"
    return fields


# ── 세계관 API ─────────────────────────────────────────────────────────────

@app.route("/api/projects/<int:project_id>/worldbuilding", methods=["GET"])
def list_worldbuilding(project_id):
    wbs = db.get_worldbuilding_by_project(project_id)
    return jsonify([w.to_dict() for w in wbs])


@app.route("/api/projects/<int:project_id>/worldbuilding", methods=["POST"])
def create_worldbuilding(project_id):
    data = request.json or {}
    wb = WorldBuilding(
        project_id=project_id,
        category=data.get("category", "기타"),
        title=data.get("title", ""),
        content=data.get("content", ""),
        source=data.get("source", "user"),
        tags=data.get("tags", ""),
    )
    wb.id = db.create_worldbuilding(wb)
    return jsonify(wb.to_dict()), 201


@app.route("/api/projects/<int:project_id>/worldbuilding/<int:wb_id>", methods=["PUT"])
def update_worldbuilding(project_id, wb_id):
    wb = db.get_worldbuilding(wb_id)
    if not wb or wb.project_id != project_id:
        abort(404)
    data = request.json or {}
    for field in ("category", "title", "content", "source", "tags"):
        if field in data:
            setattr(wb, field, data[field])
    db.update_worldbuilding(wb)
    return jsonify(wb.to_dict())


@app.route("/api/projects/<int:project_id>/worldbuilding/<int:wb_id>", methods=["DELETE"])
def delete_worldbuilding(project_id, wb_id):
    db.delete_worldbuilding(wb_id)
    return jsonify({"ok": True})


@app.route("/api/projects/<int:project_id>/worldbuilding/generate", methods=["POST"])
def generate_worldbuilding(project_id):
    p = db.get_project(project_id)
    if not p:
        abort(404)
    data = request.json or {}
    category = data.get("category", "지리/지명")
    try:
        ctx = _make_context_manager()
        raw = ctx.generate_worldbuilding(project_id, category)
        return jsonify({"category": category, "content": raw})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── AI 생성 API ────────────────────────────────────────────────────────────

@app.route("/api/generation/plot", methods=["POST"])
def generate_plot():
    data = request.json or {}
    project_id = data.get("project_id")
    if not project_id:
        return jsonify({"error": "project_id required"}), 400
    try:
        ctx = _make_context_manager()
        plot = ctx.generate_plot(project_id)
        return jsonify({"plot": plot})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/generation/single", methods=["POST"])
def generate_single():
    """단일 챕터 생성"""
    data = request.json or {}
    project_id = data.get("project_id")
    chapter_num = int(data.get("chapter_num", 1))
    user_instruction = data.get("instruction", "")
    target_length = int(data.get("target_length", 2000))

    p = db.get_project(project_id)
    if not p:
        return jsonify({"error": "Project not found"}), 404

    # 목차에서 제목 조회
    outlines = db.get_outlines_by_project(project_id)
    outline_map = {o.chapter_num: o for o in outlines}
    outline = outline_map.get(chapter_num)
    chapter_title = outline.title if outline else f"{chapter_num}편"

    try:
        ctx = _make_context_manager()
        t0 = time.time()
        content = ctx.generate_chapter(
            project_id=project_id,
            user_request=user_instruction or "자연스러운 다음 내용을 작성해주세요.",
            target_length=target_length,
            chapter_num=chapter_num,
            chapter_title=chapter_title,
        )
        elapsed = time.time() - t0
        word_count = len(content)

        existing = db.get_chapter_by_num(project_id, chapter_num)
        if existing:
            db.save_generation_history(existing.id, existing.content)
            existing.content = content
            existing.title = chapter_title
            existing.word_count = word_count
            existing.status = "generated"
            existing.generation_time = elapsed
            db.update_chapter(existing)
            chapter = existing
        else:
            chapter = Chapter(
                project_id=project_id,
                outline_id=outline.id if outline else None,
                chapter_num=chapter_num,
                title=chapter_title,
                content=content,
                word_count=word_count,
                status="generated",
                generation_time=elapsed,
            )
            chapter.id = db.create_chapter(chapter)

        # 요약 생성
        try:
            ch = db.get_chapter(chapter.id)
            if ch:
                ctx.auto_summarize_and_save(ch)
                chapter = db.get_chapter(chapter.id)
        except Exception:
            pass

        return jsonify(chapter.to_dict())
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/generation/batch", methods=["POST"])
def generate_batch():
    """배치/전체 자동 생성 작업 시작"""
    data = request.json or {}
    project_id = data.get("project_id")
    job_type = data.get("job_type", "batch")  # 'batch' | 'full_auto'
    start_chapter = int(data.get("start_chapter", 1))
    end_chapter = int(data.get("end_chapter", start_chapter))
    target_length = int(data.get("target_length", 2000))
    user_instruction = data.get("instruction", "")

    p = db.get_project(project_id)
    if not p:
        return jsonify({"error": "Project not found"}), 404

    # 기존 활성 작업 확인
    active = db.get_active_job_by_project(project_id)
    if active:
        return jsonify({"error": "이미 실행 중인 생성 작업이 있습니다.", "job": active.to_dict()}), 409

    job = GenerationJob(
        project_id=project_id,
        job_type=job_type,
        start_chapter=start_chapter,
        end_chapter=end_chapter,
        total_chapters=end_chapter - start_chapter + 1,
        status="pending",
    )
    job.id = db.create_generation_job(job)

    ctx = _make_context_manager()
    generator = AutoGenerator(db, ctx, on_progress=_on_job_progress)
    generator.start_job(job, target_length=target_length, user_instruction=user_instruction)

    return jsonify(job.to_dict()), 202


@app.route("/api/generation/jobs/<int:job_id>", methods=["GET"])
def get_job(job_id):
    job = db.get_generation_job(job_id)
    if not job:
        abort(404)
    return jsonify(job.to_dict())


@app.route("/api/generation/jobs/<int:job_id>/pause", methods=["POST"])
def pause_generation(job_id):
    pause_job(job_id)
    job = db.get_generation_job(job_id)
    if job:
        job.status = "paused"
        db.update_generation_job(job)
    return jsonify({"ok": True})


@app.route("/api/generation/jobs/<int:job_id>/resume", methods=["POST"])
def resume_generation(job_id):
    resume_job(job_id)
    job = db.get_generation_job(job_id)
    if job:
        job.status = "running"
        db.update_generation_job(job)
    return jsonify({"ok": True})


@app.route("/api/generation/jobs/<int:job_id>/cancel", methods=["POST"])
def cancel_generation(job_id):
    cancel_job(job_id)
    job = db.get_generation_job(job_id)
    if job:
        job.status = "failed"
        job.error_message = "사용자에 의해 취소됨"
        db.update_generation_job(job)
    return jsonify({"ok": True})


@app.route("/api/generation/summary", methods=["POST"])
def generate_summary():
    data = request.json or {}
    chapter_id = data.get("chapter_id")
    ch = db.get_chapter(chapter_id)
    if not ch:
        return jsonify({"error": "Chapter not found"}), 404
    try:
        ctx = _make_context_manager()
        summary = ctx.auto_summarize_and_save(ch)
        return jsonify({"summary": summary})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── 내보내기 API ───────────────────────────────────────────────────────────

@app.route("/api/export/<int:project_id>/txt", methods=["GET"])
def export_txt(project_id):
    p = db.get_project(project_id)
    if not p:
        abort(404)
    chapters = db.get_chapters_by_project(project_id)

    buf = io.StringIO()
    buf.write(f"{'=' * 50}\n  {p.title}\n{'=' * 50}\n\n")
    for ch in sorted(chapters, key=lambda c: c.chapter_num):
        title_str = f": {ch.title}" if ch.title else ""
        buf.write(f"\n=== {ch.chapter_num}편{title_str} ===\n\n")
        buf.write(ch.content or "")
        buf.write("\n\n")

    content = buf.getvalue()
    return send_file(
        io.BytesIO(content.encode("utf-8")),
        mimetype="text/plain; charset=utf-8",
        as_attachment=True,
        download_name=f"{p.title}.txt",
    )


@app.route("/api/export/<int:project_id>/docx", methods=["GET"])
def export_docx(project_id):
    p = db.get_project(project_id)
    if not p:
        abort(404)
    chapters = db.get_chapters_by_project(project_id)

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return jsonify({"error": "python-docx not installed"}), 500

    doc = Document()
    title_para = doc.add_heading(p.title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for ch in sorted(chapters, key=lambda c: c.chapter_num):
        doc.add_page_break()
        title_str = f": {ch.title}" if ch.title else ""
        doc.add_heading(f"{ch.chapter_num}편{title_str}", level=1)
        if ch.content:
            for para in ch.content.split("\n"):
                if para.strip():
                    doc.add_paragraph(para)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"{p.title}.docx",
    )


# ── SocketIO 이벤트 ────────────────────────────────────────────────────────

@socketio.on("connect")
def on_connect():
    pass


@socketio.on("disconnect")
def on_disconnect():
    pass


# ── 실행 ───────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    debug = os.environ.get("FLASK_DEBUG", "false").lower() == "true"
    # allow_unsafe_werkzeug is needed for the development server with threading mode
    print(f"NovelWriter Local 시작: http://localhost:{port}")
    socketio.run(app, host="0.0.0.0", port=port, debug=debug, allow_unsafe_werkzeug=not debug)
