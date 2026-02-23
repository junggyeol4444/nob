"""내보내기 유틸리티 (txt, docx)"""
import os
from typing import List
from ..database.models import Chapter


def export_chapter_txt(chapter: Chapter, file_path: str):
    """단일 챕터를 txt 파일로 저장"""
    with open(file_path, "w", encoding="utf-8") as f:
        if chapter.title:
            f.write(f"=== {chapter.chapter_num}화: {chapter.title} ===\n\n")
        else:
            f.write(f"=== {chapter.chapter_num}화 ===\n\n")
        f.write(chapter.content or "")


def export_all_chapters_txt(project_title: str, chapters: List[Chapter], file_path: str):
    """모든 챕터를 하나의 txt 파일로 통합 저장"""
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(f"{'=' * 50}\n")
        f.write(f"  {project_title}\n")
        f.write(f"{'=' * 50}\n\n")
        for ch in sorted(chapters, key=lambda c: c.chapter_num):
            if ch.title:
                f.write(f"\n=== {ch.chapter_num}화: {ch.title} ===\n\n")
            else:
                f.write(f"\n=== {ch.chapter_num}화 ===\n\n")
            f.write(ch.content or "")
            f.write("\n\n")


def export_all_chapters_docx(project_title: str, chapters: List[Chapter], file_path: str):
    """모든 챕터를 docx 파일로 내보내기"""
    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise ImportError(
            "python-docx가 설치되지 않았습니다. pip install python-docx 를 실행하세요."
        )

    doc = Document()

    # 제목 설정
    title_para = doc.add_heading(project_title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for ch in sorted(chapters, key=lambda c: c.chapter_num):
        doc.add_page_break()
        if ch.title:
            heading_text = f"{ch.chapter_num}화: {ch.title}"
        else:
            heading_text = f"{ch.chapter_num}화"
        doc.add_heading(heading_text, level=1)

        if ch.content:
            for paragraph in ch.content.split("\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph)

    doc.save(file_path)
